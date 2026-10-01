from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Form, status
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_
from typing import List, Optional
from pydantic import BaseModel
from google import genai
from datetime import datetime
import os
import shutil
import json
import tempfile

from backend.database.database import get_db
from backend.models import models
from backend.routes.auth import get_current_user
from backend.routes.p2p_signaling import manager
from backend.services.encryption_service import encrypt_secret, decrypt_secret

router = APIRouter()

@router.get("/companies")
def get_companies(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    return db.query(models.Company).filter(
        models.Company.user_id == current_user.id
    ).order_by(models.Company.id.desc()).all()

@router.post("/companies")
def create_company(
    name: str, 
    description: str = None, 
    rate_detection: float = None,
    rate_conformation: float = None,
    rate_pose_texte: float = None,
    rate_chantant: float = None,
    supplier_email: str = None,
    target_software: str = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    db_company = models.Company(
        user_id=current_user.id,
        name=name, 
        description=description,
        rate_detection=rate_detection,
        rate_conformation=rate_conformation,
        rate_pose_texte=rate_pose_texte,
        rate_chantant=rate_chantant,
        supplier_email=supplier_email,
        target_software=target_software
    )
    db.add(db_company)
    db.commit()
    db.refresh(db_company)
    return db_company

@router.get("/companies/{company_id}")
def get_company(
    company_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    company = db.query(models.Company).filter(
        models.Company.id == company_id,
        models.Company.user_id == current_user.id
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="Company not found")
    return company

@router.put("/companies/{company_id}")
def update_company(
    company_id: int,
    name: str = None,
    description: str = None,
    rate_detection: float = None,
    rate_conformation: float = None,
    rate_pose_texte: float = None,
    rate_chantant: float = None,
    supplier_email: str = None,
    target_software: str = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    company = db.query(models.Company).filter(
        models.Company.id == company_id,
        models.Company.user_id == current_user.id
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="Company not found")
        
    if name is not None: company.name = name
    if description is not None: company.description = description
    if rate_detection is not None: company.rate_detection = rate_detection
    if rate_conformation is not None: company.rate_conformation = rate_conformation
    if rate_pose_texte is not None: company.rate_pose_texte = rate_pose_texte
    if rate_chantant is not None: company.rate_chantant = rate_chantant
    if supplier_email is not None: company.supplier_email = supplier_email
    if target_software is not None: company.target_software = target_software
    
    db.commit()
    db.refresh(company)
    return company

@router.delete("/companies/{company_id}")
def delete_company(
    company_id: int, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    company = db.query(models.Company).filter(
        models.Company.id == company_id,
        models.Company.user_id == current_user.id
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="Company not found")
        
    db.delete(company)
    db.commit()
    return {"message": "Company deleted successfully"}

# --- User Search Endpoint ---

@router.get("/users/search")
def search_users(
    q: str = "",
    limit: int = 15,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    """
    Search registered users by first name, last name, full name, or email.
    Used by the Staff management page to find user accounts and add/link them as staff.
    """
    query = q.strip().lower()
    users_query = db.query(models.User)
    
    if query:
        tokens = query.split()
        if len(tokens) >= 2:
            first_part, second_part = tokens[0], tokens[1]
            users_query = users_query.filter(
                (
                    models.User.first_name.ilike(f"%{first_part}%") & 
                    models.User.last_name.ilike(f"%{second_part}%")
                ) | (
                    models.User.first_name.ilike(f"%{second_part}%") & 
                    models.User.last_name.ilike(f"%{first_part}%")
                ) | (
                    models.User.email.ilike(f"%{query}%")
                )
            )
        else:
            term = f"%{query}%"
            users_query = users_query.filter(
                models.User.first_name.ilike(term) |
                models.User.last_name.ilike(term) |
                models.User.email.ilike(term)
            )
    
    users = users_query.limit(limit).all()
    
    return [
        {
            "id": u.id,
            "first_name": u.first_name,
            "last_name": u.last_name,
            "name": f"{u.first_name} {u.last_name}".strip(),
            "email": u.email,
            "phone_number": u.phone_number
        }
        for u in users
    ]

# --- Staff Endpoints ---

@router.get("/staff/online-status")
def get_staff_online_status():
    return {
        "online_user_ids": list(manager.get_online_user_ids())
    }

@router.get("/staff")
def get_staff(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    staff_list = db.query(models.Staff).filter(
        models.Staff.user_id == current_user.id
    ).order_by(models.Staff.id.desc()).all()

    online_user_ids = manager.get_online_user_ids()
    result = []
    for s in staff_list:
        is_online = False
        if s.staff_user_id and s.staff_user_id in online_user_ids:
            is_online = True
        elif s.email:
            linked_user = db.query(models.User).filter(models.User.email == s.email).first()
            if linked_user and linked_user.id in online_user_ids:
                is_online = True
                if not s.staff_user_id:
                    s.staff_user_id = linked_user.id
                    db.commit()

        result.append({
            "id": s.id,
            "name": s.name,
            "email": s.email,
            "task": s.task,
            "access_level": s.access_level or "spectator",
            "auto_accept_transfers": bool(s.auto_accept_transfers),
            "staff_user_id": s.staff_user_id,
            "is_online": is_online,
            "created_at": s.created_at.isoformat() if s.created_at else None
        })
    return result

@router.post("/staff")
async def create_staff(
    request: Request,
    name: str = None,
    email: str = None,
    task: str = None,
    access_level: str = None,
    auto_accept_transfers: bool = None,
    staff_user_id: int = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    try:
        body = await request.json()
    except Exception:
        body = {}

    if isinstance(body, dict):
        if name is None and "name" in body: name = body["name"]
        if email is None and "email" in body: email = body["email"]
        if task is None and "task" in body: task = body["task"]
        if access_level is None and "access_level" in body: access_level = body["access_level"]
        if auto_accept_transfers is None and "auto_accept_transfers" in body: auto_accept_transfers = body["auto_accept_transfers"]
        if staff_user_id is None and "staff_user_id" in body: staff_user_id = body["staff_user_id"]

    if not name or not str(name).strip():
        raise HTTPException(status_code=422, detail="Collaborator name is required")

    clean_access = (access_level or "spectator").strip().lower()
    if clean_access not in ["full_access", "spectator"]:
        clean_access = "spectator"

    clean_email = email.strip().lower() if email and email.strip() else None
    linked_user = None

    if staff_user_id and int(staff_user_id) > 0:
        linked_user = db.query(models.User).filter(models.User.id == int(staff_user_id)).first()
    elif clean_email:
        linked_user = db.query(models.User).filter(models.User.email == clean_email).first()

    if not linked_user:
        raise HTTPException(
            status_code=404,
            detail="Only registered users can be added as collaborators. Please ask the user to register an account first."
        )

    if linked_user.id == current_user.id:
        raise HTTPException(
            status_code=400,
            detail="You cannot add your own account as a collaborator."
        )

    # Check for duplicate
    existing_staff = db.query(models.Staff).filter(
        models.Staff.user_id == current_user.id,
        (models.Staff.staff_user_id == linked_user.id) | (models.Staff.email == linked_user.email)
    ).first()
    if existing_staff:
        raise HTTPException(
            status_code=400,
            detail="This user is already in your collaborators list."
        )

    valid_staff_user_id = linked_user.id
    clean_email = linked_user.email
    full_name = f"{linked_user.first_name} {linked_user.last_name}".strip() or str(name).strip()

    db_staff = models.Staff(
        user_id=current_user.id,
        staff_user_id=valid_staff_user_id,
        name=full_name,
        email=clean_email,
        task=str(task).strip() if task else None,
        access_level=clean_access,
        auto_accept_transfers=bool(auto_accept_transfers)
    )
    db.add(db_staff)
    db.commit()
    db.refresh(db_staff)

    online_ids = manager.get_online_user_ids()
    is_online = valid_staff_user_id in online_ids if valid_staff_user_id else False

    return {
        "id": db_staff.id,
        "name": db_staff.name,
        "email": db_staff.email,
        "task": db_staff.task,
        "access_level": db_staff.access_level,
        "auto_accept_transfers": db_staff.auto_accept_transfers,
        "staff_user_id": db_staff.staff_user_id,
        "is_online": is_online,
        "created_at": db_staff.created_at.isoformat() if db_staff.created_at else None
    }

@router.get("/staff/{staff_id}")
def get_staff_member(
    staff_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    staff_member = db.query(models.Staff).filter(
        models.Staff.id == staff_id,
        models.Staff.user_id == current_user.id
    ).first()
    if not staff_member:
        raise HTTPException(status_code=404, detail="Staff member not found")

    online_ids = manager.get_online_user_ids()
    is_online = staff_member.staff_user_id in online_ids if staff_member.staff_user_id else False

    return {
        "id": staff_member.id,
        "name": staff_member.name,
        "email": staff_member.email,
        "task": staff_member.task,
        "access_level": staff_member.access_level or "spectator",
        "auto_accept_transfers": bool(staff_member.auto_accept_transfers),
        "staff_user_id": staff_member.staff_user_id,
        "is_online": is_online,
        "created_at": staff_member.created_at.isoformat() if staff_member.created_at else None
    }

@router.put("/staff/{staff_id}")
async def update_staff(
    staff_id: int,
    request: Request,
    name: str = None,
    email: str = None,
    task: str = None,
    access_level: str = None,
    auto_accept_transfers: bool = None,
    staff_user_id: int = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    try:
        body = await request.json()
    except Exception:
        body = {}

    if isinstance(body, dict):
        if name is None and "name" in body: name = body["name"]
        if email is None and "email" in body: email = body["email"]
        if task is None and "task" in body: task = body["task"]
        if access_level is None and "access_level" in body: access_level = body["access_level"]
        if auto_accept_transfers is None and "auto_accept_transfers" in body: auto_accept_transfers = body["auto_accept_transfers"]
        if staff_user_id is None and "staff_user_id" in body: staff_user_id = body["staff_user_id"]

    staff = db.query(models.Staff).filter(
        models.Staff.id == staff_id,
        models.Staff.user_id == current_user.id
    ).first()
    if not staff:
        raise HTTPException(status_code=404, detail="Staff not found")

    if name is not None: staff.name = str(name).strip()
    if email is not None:
        clean_email = str(email).strip().lower() if str(email).strip() else None
        staff.email = clean_email
        if clean_email and not staff.staff_user_id:
            linked_user = db.query(models.User).filter(models.User.email == clean_email).first()
            if linked_user:
                staff.staff_user_id = linked_user.id
    if task is not None: staff.task = str(task).strip() if task else None
    if access_level is not None:
        clean_level = str(access_level).strip().lower()
        if clean_level in ["full_access", "spectator"]:
            staff.access_level = clean_level
    if auto_accept_transfers is not None:
        staff.auto_accept_transfers = bool(auto_accept_transfers)
    if staff_user_id is not None:
        if int(staff_user_id) > 0:
            linked_user = db.query(models.User).filter(models.User.id == int(staff_user_id)).first()
            staff.staff_user_id = linked_user.id if linked_user else None
        else:
            staff.staff_user_id = None

    db.commit()
    db.refresh(staff)

    online_ids = manager.get_online_user_ids()
    is_online = staff.staff_user_id in online_ids if staff.staff_user_id else False

    return {
        "id": staff.id,
        "name": staff.name,
        "email": staff.email,
        "task": staff.task,
        "access_level": staff.access_level,
        "auto_accept_transfers": staff.auto_accept_transfers,
        "staff_user_id": staff.staff_user_id,
        "is_online": is_online,
        "created_at": staff.created_at.isoformat() if staff.created_at else None
    }

@router.patch("/staff/{staff_id}/auto-accept")
def toggle_staff_auto_accept(
    staff_id: int,
    auto_accept: Optional[bool] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    staff = db.query(models.Staff).filter(
        models.Staff.id == staff_id,
        models.Staff.user_id == current_user.id
    ).first()
    if not staff:
        raise HTTPException(status_code=404, detail="Staff not found")

    if auto_accept is None:
        staff.auto_accept_transfers = not staff.auto_accept_transfers
    else:
        staff.auto_accept_transfers = bool(auto_accept)

    db.commit()
    db.refresh(staff)
    return {
        "id": staff.id,
        "auto_accept_transfers": staff.auto_accept_transfers,
        "message": f"Auto-accept transfers {'enabled' if staff.auto_accept_transfers else 'disabled'} successfully."
    }

@router.delete("/staff/{staff_id}")
def delete_staff(
    staff_id: int, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    staff = db.query(models.Staff).filter(
        models.Staff.id == staff_id,
        models.Staff.user_id == current_user.id
    ).first()
    if not staff:
        raise HTTPException(status_code=404, detail="Staff not found")
        
    db.delete(staff)
    db.commit()
    return {"message": "Staff deleted successfully"}

# --- Project Endpoints (User-Scoped) ---

@router.get("/projects")
def get_projects(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    owned_projects = db.query(models.Project).filter(
        models.Project.user_id == current_user.id
    ).order_by(models.Project.id.desc()).all()

    # Find shared projects where current user is linked as a collaborator
    collaborations = db.query(models.Staff).filter(
        (models.Staff.staff_user_id == current_user.id) | (models.Staff.email == current_user.email)
    ).all()

    seen_ids = {p.id for p in owned_projects}
    results = []

    for p in owned_projects:
        results.append({
            "id": p.id,
            "user_id": p.user_id,
            "name": p.name,
            "company_id": p.company_id,
            "company_name": p.company_name,
            "folder_path": p.folder_path,
            "target_software": p.target_software,
            "project_type": p.project_type,
            "deadline": p.deadline.isoformat() if p.deadline else None,
            "total_time": p.total_time,
            "status": p.status,
            "is_owner": True,
            "access_level": "full_access",
            "owner_id": current_user.id,
            "owner_name": f"{current_user.first_name} {current_user.last_name}".strip(),
            "created_at": p.created_at.isoformat() if p.created_at else None
        })

    for collab in collaborations:
        if collab.user_id and collab.user_id != current_user.id:
            owner = db.query(models.User).filter(models.User.id == collab.user_id).first()
            shared_list = db.query(models.Project).filter(models.Project.user_id == collab.user_id).all()
            for sp in shared_list:
                if sp.id not in seen_ids:
                    seen_ids.add(sp.id)
                    results.append({
                        "id": sp.id,
                        "user_id": sp.user_id,
                        "name": sp.name,
                        "company_id": sp.company_id,
                        "company_name": sp.company_name,
                        "folder_path": sp.folder_path,
                        "target_software": sp.target_software,
                        "project_type": sp.project_type,
                        "deadline": sp.deadline.isoformat() if sp.deadline else None,
                        "total_time": sp.total_time,
                        "status": sp.status,
                        "is_owner": False,
                        "access_level": collab.access_level or "spectator",
                        "owner_id": collab.user_id,
                        "owner_name": f"{owner.first_name} {owner.last_name}".strip() if owner else "Owner",
                        "created_at": sp.created_at.isoformat() if sp.created_at else None
                    })

    return results

@router.get("/browse-folder")
def browse_folder():
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        folder_path = filedialog.askdirectory(parent=root, title="Select Project Folder")
        root.destroy()
        return {"path": folder_path}
    except Exception as e:
        print(f"Error opening folder dialog: {e}")
        return {"path": ""}

@router.post("/projects")
def create_project(
    name: str, 
    company_name: str = None, 
    folder_path: str = None, 
    target_software: str = None, 
    project_type: str = None,
    deadline: str = None, 
    company_id: int = None, 
    total_time: int = None,
    description: str = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    dl_dt = None
    if deadline:
        try:
            dl_dt = datetime.fromisoformat(deadline.replace("Z", "+00:00"))
        except:
            pass

    verified_company_id = None
    verified_company_name = company_name
    if company_id:
        company = db.query(models.Company).filter(
            models.Company.id == company_id,
            models.Company.user_id == current_user.id
        ).first()
        if not company:
            raise HTTPException(status_code=400, detail="Invalid company selected or company does not belong to you.")
        verified_company_id = company.id
        if not verified_company_name:
            verified_company_name = company.name

    db_project = models.Project(
        user_id=current_user.id,
        name=name, 
        company_id=verified_company_id,
        company_name=verified_company_name,
        folder_path=folder_path,
        target_software=target_software,
        project_type=project_type,
        deadline=dl_dt,
        total_time=total_time,
        description=description
    )
    db.add(db_project)
    db.commit()
    db.refresh(db_project)
    return db_project

@router.get("/projects/{project_id}")
def get_project(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    is_owner = (project.user_id == current_user.id)
    access_level = "full_access" if is_owner else None
    owner = None

    if not is_owner:
        collab = db.query(models.Staff).filter(
            models.Staff.user_id == project.user_id,
            or_(
                models.Staff.staff_user_id == current_user.id,
                models.Staff.email == current_user.email
            )
        ).first()
        if not collab:
            raise HTTPException(status_code=404, detail="Project not found")
        access_level = collab.access_level or "spectator"
        owner = db.query(models.User).filter(models.User.id == project.user_id).first()

    return {
        "id": project.id,
        "user_id": project.user_id,
        "name": project.name,
        "company_id": project.company_id,
        "company_name": project.company_name,
        "folder_path": project.folder_path,
        "target_software": project.target_software,
        "project_type": project.project_type,
        "deadline": project.deadline.isoformat() if project.deadline else None,
        "total_time": project.total_time,
        "status": project.status,
        "is_owner": is_owner,
        "access_level": access_level,
        "owner_id": project.user_id,
        "owner_name": f"{owner.first_name} {owner.last_name}".strip() if owner else "Owner",
        "created_at": project.created_at.isoformat() if project.created_at else None
    }

@router.put("/projects/{project_id}")
def update_project(
    project_id: int,
    name: str = None, 
    company_name: str = None, 
    folder_path: str = None, 
    target_software: str = None, 
    project_type: str = None,
    deadline: str = None, 
    company_id: int = None,
    total_time: int = None,
    description: str = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    project = db.query(models.Project).filter(
        models.Project.id == project_id,
        models.Project.user_id == current_user.id
    ).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    if company_id is not None:
        if company_id > 0:
            comp = db.query(models.Company).filter(
                models.Company.id == company_id,
                models.Company.user_id == current_user.id
            ).first()
            if not comp:
                raise HTTPException(status_code=400, detail="Invalid company selected or company does not belong to you.")
            project.company_id = comp.id
            if company_name is None:
                project.company_name = comp.name
        else:
            project.company_id = None
        
    if name is not None: project.name = name
    if description is not None: project.description = description
    if company_name is not None: project.company_name = company_name
    if folder_path is not None: project.folder_path = folder_path
    if target_software is not None: project.target_software = target_software
    if project_type is not None: project.project_type = project_type
    if total_time is not None: project.total_time = total_time
    if deadline is not None:
        try:
            project.deadline = datetime.fromisoformat(deadline.replace("Z", "+00:00"))
        except:
            pass
            
    db.commit()
    db.refresh(project)
    return project

@router.delete("/projects/{project_id}")
def delete_project(
    project_id: int, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    project = db.query(models.Project).filter(
        models.Project.id == project_id,
        models.Project.user_id == current_user.id
    ).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
        
    # Clean up project uploaded files if any
    upload_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "uploads", "projects", str(project.id))
    if os.path.exists(upload_dir):
        try:
            shutil.rmtree(upload_dir)
        except Exception as e:
            print(f"Error removing project upload directory: {e}")

    db.delete(project)
    db.commit()
    return {"message": "Project deleted successfully"}

class ProjectIndexPayload(BaseModel):
    files: List[dict]
    folder_path: Optional[str] = None

@router.get("/projects/{project_id}/files")
def get_project_files(
    project_id: int, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    is_owner = (project.user_id == current_user.id)
    access_level = "full_access" if is_owner else None

    if not is_owner:
        collab = db.query(models.Staff).filter(
            models.Staff.user_id == project.user_id,
            or_(
                models.Staff.staff_user_id == current_user.id,
                models.Staff.email == current_user.email
            )
        ).first()
        if not collab:
            raise HTTPException(status_code=404, detail="Project not found")
        access_level = collab.access_level or "spectator"

    files_list = []
    seen_filenames = set()

    # 1. Load saved indexed files metadata if present
    if project.files_index:
        try:
            indexed = json.loads(project.files_index)
            if isinstance(indexed, list):
                for item in indexed:
                    name = item.get("name")
                    if name and name not in seen_filenames:
                        files_list.append({
                            "name": name,
                            "path": item.get("path") or name,
                            "type": item.get("type", "script"),
                            "size": item.get("size", 0),
                            "is_owner": is_owner,
                            "access_level": access_level,
                            "owner_id": project.user_id
                        })
                        seen_filenames.add(name)
        except Exception as e:
            print(f"Error parsing project files_index: {e}")

    # 2. If running locally (desktop app or local server) and folder_path exists, scan direct disk
    if project.folder_path and os.path.exists(project.folder_path):
        try:
            for entry in os.scandir(project.folder_path):
                if entry.is_file() and not entry.name.startswith('.'):
                    if entry.name not in seen_filenames:
                        ext = os.path.splitext(entry.name)[1].lower()
                        is_video = ext in ['.mp4', '.mkv', '.avi', '.mov', '.wmv', '.webm']
                        is_script = ext in ['.txt', '.srt', '.doc', '.docx', '.pdf', '.rtf', '.xml', '.json', '.mos', '.awf']
                        is_audio = ext in ['.mp3', '.wav', '.m4a', '.aac', '.ogg', '.flac']
                        if is_video or is_script or is_audio:
                            try:
                                size_bytes = entry.stat().st_size
                            except Exception:
                                size_bytes = 0
                            files_list.append({
                                "name": entry.name,
                                "path": entry.path,
                                "type": "video" if is_video else ("audio" if is_audio else "script"),
                                "size": size_bytes,
                                "is_owner": is_owner,
                                "access_level": access_level,
                                "owner_id": project.user_id
                            })
                            seen_filenames.add(entry.name)
        except Exception as e:
            print(f"Local folder scan notification: {e}")

    return files_list

@router.put("/projects/{project_id}/index")
def update_project_files_index(
    project_id: int,
    payload: ProjectIndexPayload,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    is_owner = (project.user_id == current_user.id)
    if not is_owner:
        collab = db.query(models.Staff).filter(
            models.Staff.user_id == project.user_id,
            or_(
                models.Staff.staff_user_id == current_user.id,
                models.Staff.email == current_user.email
            )
        ).first()
        if not collab or collab.access_level != "full_access":
            raise HTTPException(status_code=403, detail="Permission denied. Only owners or full access collaborators can update project index.")

    if payload.folder_path is not None:
        project.folder_path = payload.folder_path.strip()

    clean_files = []
    seen = set()
    for f in payload.files:
        name = str(f.get("name", "")).strip()
        if not name or name in seen:
            continue
        ext = os.path.splitext(name)[1].lower()
        is_video = ext in ['.mp4', '.mkv', '.avi', '.mov', '.wmv', '.webm']
        is_script = ext in ['.txt', '.srt', '.doc', '.docx', '.pdf', '.rtf', '.xml', '.json', '.mos', '.awf']
        is_audio = ext in ['.mp3', '.wav', '.m4a', '.aac', '.ogg', '.flac']
        clean_files.append({
            "name": name,
            "path": f.get("path") or name,
            "type": "video" if is_video else ("audio" if is_audio else "script"),
            "size": int(f.get("size") or 0),
            "is_owner": is_owner,
            "owner_id": project.user_id,
            "access_level": "full_access" if is_owner else "spectator"
        })
        seen.add(name)

    project.files_index = json.dumps(clean_files)
    db.commit()
    db.refresh(project)
    return {"success": True, "files": clean_files, "folder_path": project.folder_path}

@router.delete("/projects/{project_id}/files")
def delete_project_file(
    project_id: int,
    filename: str,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    is_owner = (project.user_id == current_user.id)
    if not is_owner:
        raise HTTPException(status_code=403, detail="Only the project owner can remove files from index.")

    safe_name = os.path.basename(filename)
    if project.files_index:
        try:
            indexed = json.loads(project.files_index)
            if isinstance(indexed, list):
                updated_indexed = [f for f in indexed if f.get("name") != safe_name]
                project.files_index = json.dumps(updated_indexed)
                db.commit()
                return {"message": f"File {safe_name} removed from index"}
        except Exception:
            pass

    return {"message": f"File {safe_name} processed"}

@router.get("/stream-file")
def stream_file(path: str):
    import os
    resolved_path = path
    if not os.path.exists(resolved_path) or not os.path.isfile(resolved_path):
        backend_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        rel_check = os.path.join(backend_root, path.lstrip("/\\"))
        if os.path.exists(rel_check) and os.path.isfile(rel_check):
            resolved_path = rel_check
        else:
            raise HTTPException(status_code=404, detail="File not found on server")
    
    return FileResponse(resolved_path, filename=os.path.basename(resolved_path))

@router.get("/appointments")
def get_appointments(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    return db.query(models.Appointment).filter(
        models.Appointment.user_id == current_user.id
    ).order_by(models.Appointment.start_time.asc()).all()

from datetime import datetime

@router.post("/appointments")
def create_appointment(
    title: str, 
    type: str, 
    start_time: str, 
    end_time: str, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    try:
        start_dt = datetime.fromisoformat(start_time.replace("Z", "+00:00"))
        end_dt = datetime.fromisoformat(end_time.replace("Z", "+00:00"))
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date format. Use ISO format.")
        
    db_appointment = models.Appointment(
        user_id=current_user.id,
        title=title, 
        type=type,
        start_time=start_dt,
        end_time=end_dt
    )
    db.add(db_appointment)
    db.commit()
    db.refresh(db_appointment)
    return db_appointment

@router.put("/appointments/{appointment_id}")
def update_appointment(
    appointment_id: int,
    title: str = None,
    type: str = None,
    start_time: str = None,
    end_time: str = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    appointment = db.query(models.Appointment).filter(
        models.Appointment.id == appointment_id,
        models.Appointment.user_id == current_user.id
    ).first()
    if not appointment:
        raise HTTPException(status_code=404, detail="Appointment not found")

    if title is not None: appointment.title = title
    if type is not None: appointment.type = type
    if start_time is not None:
        try:
            appointment.start_time = datetime.fromisoformat(start_time.replace("Z", "+00:00"))
        except ValueError:
            pass
    if end_time is not None:
        try:
            appointment.end_time = datetime.fromisoformat(end_time.replace("Z", "+00:00"))
        except ValueError:
            pass

    db.commit()
    db.refresh(appointment)
    return appointment

@router.delete("/appointments/{appointment_id}")
def delete_appointment(
    appointment_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    appointment = db.query(models.Appointment).filter(
        models.Appointment.id == appointment_id,
        models.Appointment.user_id == current_user.id
    ).first()
    if not appointment:
        raise HTTPException(status_code=404, detail="Appointment not found")
        
    db.delete(appointment)
    db.commit()
    return {"message": "Appointment deleted successfully"}

import os
import json
import tempfile
from backend.services.script_parser import safe_validate_and_convert, align_timecodes_with_gemini
from backend.services.script_validator import ValidationStatus

ALLOWED_SCRIPT_EXTENSIONS = {".docx", ".pdf", ".txt"}
MAX_SCRIPT_SIZE = 50 * 1024 * 1024  # 50 MB

ALLOWED_MEDIA_EXTENSIONS = {".mp3", ".wav", ".mp4", ".mkv", ".mov", ".m4a", ".aac", ".avi", ".webm"}
MAX_MEDIA_SIZE = 500 * 1024 * 1024  # 500 MB

@router.post("/convert")
async def convert_script(
    file: UploadFile = File(...),
    current_user: models.User = Depends(get_current_user)
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Missing filename")
    
    suffix = os.path.splitext(file.filename)[1].lower()
    if suffix not in ALLOWED_SCRIPT_EXTENSIONS:
        raise HTTPException(
            status_code=400, 
            detail=f"Unsupported file format '{suffix}'. Allowed formats: {', '.join(ALLOWED_SCRIPT_EXTENSIONS)}"
        )

    temp_path = None
    try:
        total_size = 0
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_path = temp_file.name
            while chunk := await file.read(1024 * 1024):
                total_size += len(chunk)
                if total_size > MAX_SCRIPT_SIZE:
                    raise HTTPException(
                        status_code=413,
                        detail="Script file exceeds the maximum allowed size (50 MB)."
                    )
                temp_file.write(chunk)
        
        # Process the file using user's keys
        report, raw_rows, format_a_cues = safe_validate_and_convert(temp_path, user_id=current_user.id)
        
        if report.status == ValidationStatus.INVALID:
            raise HTTPException(
                status_code=400, 
                detail={"title": report.user_title, "message": report.user_message, "suggestion": getattr(report, 'suggestion', '')}
            )
            
        return {
            "success": True,
            "filename": file.filename,
            "raw_rows": raw_rows,
            "cues": format_a_cues
        }
    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass

@router.post("/align")
async def align_script(
    cues: str = Form(...), 
    start_tc: str = Form(...), 
    media: UploadFile = File(...),
    current_user: models.User = Depends(get_current_user)
):
    if not media.filename:
        raise HTTPException(status_code=400, detail="Missing media filename")
        
    suffix = os.path.splitext(media.filename)[1].lower()
    if suffix not in ALLOWED_MEDIA_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported media format '{suffix}'. Allowed formats: {', '.join(ALLOWED_MEDIA_EXTENSIONS)}"
        )

    temp_media_path = None
    try:
        total_size = 0
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_media:
            temp_media_path = temp_media.name
            while chunk := await media.read(1024 * 1024):
                total_size += len(chunk)
                if total_size > MAX_MEDIA_SIZE:
                    raise HTTPException(
                        status_code=413,
                        detail="Media file exceeds the maximum allowed size (500 MB)."
                    )
                temp_media.write(chunk)
            
        parsed_cues = json.loads(cues)
        generator = align_timecodes_with_gemini(parsed_cues, temp_media_path, start_tc, user_id=current_user.id)
        aligned_cues = list(generator)
            
        return {"success": True, "aligned_cues": aligned_cues}
    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if temp_media_path and os.path.exists(temp_media_path):
            try:
                os.remove(temp_media_path)
            except Exception:
                pass

def create_mosaic_excel(parsed_cues, filename: str = None):
    import openpyxl
    from openpyxl.styles import Font, Alignment
    from datetime import datetime
    import re
    
    wb = openpyxl.Workbook()
    
    # 1. Sheet: Dialogue List
    ws_dialogue = wb.active
    ws_dialogue.title = "Dialogue List"
    
    header_font = Font(name="Calibri", size=13, bold=True)
    data_font = Font(name="Calibri", size=11, bold=False)
    summary_header_font = Font(name="Calibri", size=13, bold=True)
    summary_total_font = Font(name="Calibri", size=11, bold=True)
    project_label_font = Font(name="Calibri", size=13, bold=True)
    project_val_font = Font(name="Calibri", size=11, bold=False)
    
    dialogue_headers = [
        ('IN-TIMECODE', 31.29),
        ('OUT-TIMECODE', 13.0),
        ('SOURCE', 13.0),
        ('TRANSCRIPTION', 93.71),
        ('ANNOTATIONS', 13.0),
        ('TAGS', 31.29),
        ('SOURCE ONSCREEN', 13.0),
        ('FN TREATMENT', 13.0),
        ('FN POSITION', 13.0)
    ]
    
    for col_idx, (hdr, width) in enumerate(dialogue_headers, start=1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws_dialogue.column_dimensions[col_letter].width = width
        cell = ws_dialogue.cell(row=1, column=col_idx, value=hdr)
        cell.font = header_font
        
    char_counts = {}
    last_out_tc = ""
    
    for row_idx, cue in enumerate(parsed_cues, start=2):
        in_tc = str(cue.get("in", "") or "").strip()
        out_tc = str(cue.get("out", "") or "").strip()
        if out_tc:
            last_out_tc = out_tc
        elif in_tc and not last_out_tc:
            last_out_tc = in_tc
            
        character = str(cue.get("character", "") or "").strip()
        dialogue = str(cue.get("dialogue", "") or "")
        annotations = str(cue.get("annotations", "") or "").strip()
        tags = str(cue.get("tags", "") or "").strip()
        source_onscreen = str(cue.get("source_onscreen", "") or "").strip()
        if not source_onscreen and character:
            source_onscreen = "ON"
        fn_treatment = str(cue.get("fn_treatment", "") or "").strip()
        fn_position = str(cue.get("fn_position", "") or "bottom").strip()
        if not fn_position:
            fn_position = "bottom"
            
        row_vals = [
            in_tc,
            out_tc,
            character,
            dialogue,
            annotations,
            tags,
            source_onscreen,
            fn_treatment,
            fn_position
        ]
        
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_dialogue.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            if col_idx == 4:  # TRANSCRIPTION
                cell.alignment = Alignment(wrap_text=True)
                
        # Word counts tracking
        if character:
            if character not in char_counts:
                char_counts[character] = {
                    "dialog": 0.0,
                    "transcription": 0.0,
                    "foreign": 0.0,
                    "music": 0.0,
                    "burnedin": 0.0,
                    "onscreen": 0.0,
                    "total": 0.0
                }
            words = float(len(re.findall(r'\b\w+\b', dialogue))) if dialogue else 0.0
            tags_upper = tags.upper()
            if "SONG" in tags_upper or "MUSIC" in tags_upper:
                char_counts[character]["music"] += words
            elif "FOREIGN" in tags_upper:
                char_counts[character]["foreign"] += words
            elif "BURNEDIN" in tags_upper or "SUBTITLE" in tags_upper:
                char_counts[character]["burnedin"] += words
            elif "ONSCREEN" in tags_upper or "TITLE" in tags_upper:
                char_counts[character]["onscreen"] += words
            else:
                char_counts[character]["dialog"] += words
            
            if dialogue.strip().startswith("[") and dialogue.strip().endswith("]"):
                char_counts[character]["transcription"] += words
                
            char_counts[character]["total"] += words

    # 2. Sheet: Word Count Summary
    ws_wc = wb.create_sheet(title="Word Count Summary")
    wc_headers = [
        ('CHARACTER NAME', 62.43),
        ('DIALOG WORD COUNT', 28.14),
        ('TRANSCRIPTION WORD COUNT', 34.43),
        ('FOREIGN DIALOGUE WORD COUNT (based on tag)', 45.57),
        ('MUSIC AND SONG WORD COUNT (based on tag)', 45.0),
        ('BURNEDIN SUBTITLE TEXT WORD COUNT (based on tags)', 37.43),
        ('ONSCREEN TEXT WORD COUNT (based on tags)', 30.0),
        ('TOTAL WORD COUNT BY CHARACTER', 34.43)
    ]
    for col_idx, (hdr, width) in enumerate(wc_headers, start=1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws_wc.column_dimensions[col_letter].width = width
        cell = ws_wc.cell(row=1, column=col_idx, value=hdr)
        cell.font = summary_header_font

    current_wc_row = 2
    cat_totals = [0.0] * 7
    for char, counts in sorted(char_counts.items()):
        ws_wc.cell(row=current_wc_row, column=1, value=char).font = data_font
        vals = [
            counts["dialog"],
            counts["transcription"],
            counts["foreign"],
            counts["music"],
            counts["burnedin"],
            counts["onscreen"],
            counts["total"]
        ]
        for c_idx, val in enumerate(vals, start=2):
            cell = ws_wc.cell(row=current_wc_row, column=c_idx, value=val)
            cell.font = data_font
            cat_totals[c_idx - 2] += val
        current_wc_row += 1

    # Total row
    ws_wc.cell(row=current_wc_row, column=1, value="TOTAL WORD COUNT BY TEXT CATEGORY").font = summary_total_font
    for c_idx, tot in enumerate(cat_totals, start=2):
        cell = ws_wc.cell(row=current_wc_row, column=c_idx, value=tot)
        cell.font = summary_total_font

    # 3. Sheet: Project Info
    ws_proj = wb.create_sheet(title="Project Info")
    ws_proj.column_dimensions['A'].width = 33.14
    ws_proj.column_dimensions['B'].width = 62.43
    
    clean_title = os.path.splitext(os.path.basename(filename or "Project Script"))[0] if filename else "Project Script"
    runtime_mins = "10"
    if last_out_tc and ":" in last_out_tc:
        parts = last_out_tc.split(":")
        try:
            hrs = int(parts[0])
            mins = int(parts[1])
            tot_mins = hrs * 60 + mins
            runtime_mins = str(tot_mins if tot_mins > 0 else 1)
        except Exception:
            runtime_mins = "10"
            
    proj_rows = [
        ("SHOW TITLE", clean_title),
        ("EPISODE TITLE", clean_title),
        ("SEASON #", ""),
        ("EPISODE #", ""),
        ("MOVIE ID", ""),
        ("TOTAL RUNTIME", runtime_mins),
        ("VERSION", "V1.0"),
        ("PROXY ID", ""),
        ("PROXY TYPE", "FINAL_CUT"),
        ("IS QC'd", "No"),
        ("DATE", datetime.now().strftime("%A %d %B %Y %H:%M:%S"))
    ]
    
    for r_idx, (k, v) in enumerate(proj_rows, start=1):
        cell_k = ws_proj.cell(row=r_idx, column=1, value=k)
        cell_k.font = project_label_font
        cell_v = ws_proj.cell(row=r_idx, column=2, value=v)
        cell_v.font = project_val_font

    return wb

def get_user_downloads_dir() -> str:
    if os.name == 'nt':
        import winreg
        try:
            sub_key = r'SOFTWARE\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders'
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, sub_key) as key:
                val, _ = winreg.QueryValueEx(key, '{374DE290-123F-4565-9164-39C4925E467B}')
                expanded = os.path.expandvars(val)
                if os.path.exists(expanded):
                    return expanded
        except Exception:
            pass
    dl = os.path.join(os.path.expanduser("~"), "Downloads")
    os.makedirs(dl, exist_ok=True)
    return dl

def get_unique_filepath(folder: str, filename: str) -> str:
    base, ext = os.path.splitext(filename)
    candidate = os.path.join(folder, filename)
    counter = 1
    while os.path.exists(candidate):
        candidate = os.path.join(folder, f"{base} ({counter}){ext}")
        counter += 1
    return candidate

@router.post("/download")
async def download_script(
    cues: str = Form(...),
    export_mode: str = Form(...),
    filename: str = Form(None),
    save_as: bool = Form(False)
):
    try:
        parsed_cues = json.loads(cues)
        clean_name = os.path.splitext(os.path.basename(filename or "Script"))[0] if filename else "Script"
        import re
        import subprocess
        safe_base_name = re.sub(r'[^\w\-_.]', '_', clean_name)
        if not safe_base_name:
            safe_base_name = "Script"

        is_mosaic = (export_mode.strip().lower() == "mosaic")
        ext = "xlsx" if is_mosaic else "docx"
        out_filename = f"{safe_base_name}_{'Mosaic' if is_mosaic else 'ERytmo'}.{ext}"

        if save_as:
            try:
                import tkinter as tk
                from tkinter import filedialog
                root = tk.Tk()
                root.withdraw()
                root.attributes('-topmost', True)
                file_types = [("Excel Spreadsheet", "*.xlsx")] if is_mosaic else [("Word Document", "*.docx")]
                file_types.append(("All Files", "*.*"))
                target_path = filedialog.asksaveasfilename(
                    title="Save Converted Script",
                    initialfile=out_filename,
                    defaultextension=f".{ext}",
                    filetypes=file_types
                )
                root.destroy()
                if not target_path:
                    return {"success": False, "cancelled": True}
            except Exception:
                downloads_dir = get_user_downloads_dir()
                target_path = get_unique_filepath(downloads_dir, out_filename)
        else:
            downloads_dir = get_user_downloads_dir()
            target_path = get_unique_filepath(downloads_dir, out_filename)

        if is_mosaic:
            wb = create_mosaic_excel(parsed_cues, filename=clean_name)
            wb.save(target_path)
        else:
            import docx
            out_doc = docx.Document()
            for idx, cue in enumerate(parsed_cues):
                tc_in = cue.get("in", "")
                if tc_in:
                    out_doc.add_paragraph(tc_in)
                out_doc.add_paragraph(cue.get("character", ""))
                out_doc.add_paragraph(cue.get("dialogue", ""))
                if idx < len(parsed_cues) - 1:
                    out_doc.add_paragraph("")
            out_doc.save(target_path)

        # Reveal in File Explorer automatically
        try:
            if os.name == 'nt' and os.path.exists(target_path):
                subprocess.Popen(f'explorer /select,"{os.path.normpath(target_path)}"')
        except Exception:
            pass

        return {
            "success": True,
            "filename": os.path.basename(target_path),
            "saved_path": os.path.normpath(target_path),
            "export_mode": export_mode
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

DANGEROUS_EXTENSIONS = {
    ".exe", ".bat", ".cmd", ".com", ".ps1", ".vbs", ".js", ".msi", 
    ".dll", ".scr", ".pif", ".reg", ".wsf", ".hta", ".cpl", ".jar", ".py", ".sh"
}

ALLOWED_DESKTOP_EXTENSIONS = {
    ".docx", ".xlsx", ".pdf", ".txt", ".mp4", ".mp3", ".wav", ".mkv", ".mov", ".avi", ".csv"
}

def _assert_local_desktop_only(request: Request, file_path: str):
    is_prod = (
        os.getenv("ENVIRONMENT") == "production" or
        os.getenv("RENDER") is not None
    )
    if is_prod:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Local file host interaction is disabled in cloud deployments."
        )
    client_host = request.client.host if request.client else ""
    if client_host not in ("127.0.0.1", "::1", "localhost", "testclient"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Host file execution is strictly limited to localhost desktop requests."
        )
    ext = os.path.splitext(file_path)[1].lower()
    if ext in DANGEROUS_EXTENSIONS or (ext and ext not in ALLOWED_DESKTOP_EXTENSIONS):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Opening executable or unapproved file formats is strictly prohibited."
        )

@router.post("/open-file")
async def open_file_endpoint(
    request: Request,
    path: str = Form(...),
    current_user: models.User = Depends(get_current_user)
):
    _assert_local_desktop_only(request, path)
    if os.path.exists(path):
        try:
            if os.name == 'nt':
                os.startfile(os.path.normpath(path))
            return {"success": True}
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Failed to open file: {e}")
    raise HTTPException(status_code=404, detail="File not found")

@router.post("/open-folder")
async def open_folder_endpoint(
    request: Request,
    path: str = Form(...),
    current_user: models.User = Depends(get_current_user)
):
    _assert_local_desktop_only(request, path)
    folder = path if os.path.isdir(path) else os.path.dirname(path)
    if os.path.exists(folder):
        try:
            if os.name == 'nt':
                import subprocess
                if os.path.isfile(path):
                    subprocess.Popen(f'explorer /select,"{os.path.normpath(path)}"')
                else:
                    subprocess.Popen(f'explorer "{os.path.normpath(folder)}"')
            return {"success": True}
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Failed to open folder: {e}")
    raise HTTPException(status_code=404, detail="Folder not found")

CONFIG_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "config.json")

def mask_api_key(key: str) -> str:
    if not key:
        return ""
    if len(key) > 8:
        return f"{key[:4]}...{key[-4:]}"
    return key

class ApiKeyCreatePayload(BaseModel):
    provider: str
    key: str
    label: str = None
    is_active: bool = True

class ApiKeyUpdatePayload(BaseModel):
    key: str = None
    label: str = None
    is_active: bool = None

def validate_key_with_provider(provider: str, key_str: str):
    key_str = key_str.strip()
    provider_lower = provider.lower()
    
    if provider_lower == "gemini":
        if not (key_str.startswith("AIza") or key_str.startswith("AQ.")):
            raise HTTPException(status_code=400, detail="Invalid Gemini API Key format. Must start with 'AIza' or 'AQ.'")
        try:
            client = genai.Client(api_key=key_str)
            for _ in client.models.list():
                break
        except Exception as e:
            err = str(e)
            if "400" in err or "403" in err or "API_KEY" in err:
                raise HTTPException(status_code=400, detail="Gemini API Key is invalid or unauthorized.")
            raise HTTPException(status_code=400, detail=f"Failed to verify Gemini API key: {err}")

    elif provider_lower == "openai":
        if not key_str.startswith("sk-"):
            raise HTTPException(status_code=400, detail="Invalid OpenAI API Key format. Must start with 'sk-'")
        try:
            import openai
            client = openai.OpenAI(api_key=key_str)
            client.models.list()
        except Exception as e:
            err = str(e)
            if "401" in err or "invalid" in err.lower():
                raise HTTPException(status_code=400, detail="OpenAI API Key is invalid or unauthorized.")
            raise HTTPException(status_code=400, detail=f"Failed to verify OpenAI API key: {err}")

    elif provider_lower == "groq":
        if not key_str.startswith("gsk_"):
            raise HTTPException(status_code=400, detail="Invalid Groq API Key format. Must start with 'gsk_'")
        try:
            import openai
            client = openai.OpenAI(api_key=key_str, base_url="https://api.groq.com/openai/v1")
            client.models.list()
        except Exception as e:
            err = str(e)
            if "401" in err or "invalid" in err.lower():
                raise HTTPException(status_code=400, detail="Groq API Key is invalid or unauthorized.")
            raise HTTPException(status_code=400, detail=f"Failed to verify Groq API key: {err}")
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported provider: {provider}")

@router.get("/settings/keys")
def get_all_keys(
    provider: str = None, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    query = db.query(models.ApiKey).filter(models.ApiKey.user_id == current_user.id)
    if provider:
        query = query.filter(models.ApiKey.provider == provider.lower())
    keys = query.order_by(models.ApiKey.id.asc()).all()
    
    res = []
    for k in keys:
        decrypted_val = decrypt_secret(k.key)
        res.append({
            "id": k.id,
            "provider": k.provider,
            "label": k.label or f"{k.provider.capitalize()} Key #{k.id}",
            "key": decrypted_val,
            "masked_key": mask_api_key(decrypted_val),
            "is_active": k.is_active,
            "created_at": k.created_at.isoformat() if k.created_at else None
        })
    return res

@router.post("/settings/keys")
def add_api_key(
    payload: ApiKeyCreatePayload, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    provider = payload.provider.strip().lower()
    raw_key = payload.key.strip()
    
    validate_key_with_provider(provider, raw_key)
    
    existing_count = db.query(models.ApiKey).filter(
        models.ApiKey.provider == provider,
        models.ApiKey.user_id == current_user.id
    ).count()
    label = payload.label.strip() if (payload.label and payload.label.strip()) else f"{provider.capitalize()} Key {existing_count + 1}"
    
    new_key = models.ApiKey(
        user_id=current_user.id,
        provider=provider,
        key=encrypt_secret(raw_key),
        label=label,
        is_active=payload.is_active
    )
    db.add(new_key)
    db.commit()
    db.refresh(new_key)
    
    return {
        "id": new_key.id,
        "provider": new_key.provider,
        "label": new_key.label,
        "masked_key": mask_api_key(raw_key),
        "is_active": new_key.is_active,
        "created_at": new_key.created_at.isoformat() if new_key.created_at else None
    }

@router.put("/settings/keys/{key_id}")
def update_api_key_by_id(
    key_id: int, 
    payload: ApiKeyUpdatePayload, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    db_key = db.query(models.ApiKey).filter(
        models.ApiKey.id == key_id,
        models.ApiKey.user_id == current_user.id
    ).first()
    if not db_key:
        raise HTTPException(status_code=404, detail="API Key not found")
        
    masked_display = mask_api_key(decrypt_secret(db_key.key))
    if payload.key is not None and payload.key.strip():
        new_key_str = payload.key.strip()
        validate_key_with_provider(db_key.provider, new_key_str)
        db_key.key = encrypt_secret(new_key_str)
        masked_display = mask_api_key(new_key_str)
        
    if payload.label is not None:
        db_key.label = payload.label.strip()
        
    if payload.is_active is not None:
        db_key.is_active = payload.is_active
        
    db.commit()
    db.refresh(db_key)
    return {
        "id": db_key.id,
        "provider": db_key.provider,
        "label": db_key.label,
        "masked_key": masked_display,
        "is_active": db_key.is_active,
        "created_at": db_key.created_at.isoformat() if db_key.created_at else None
    }

@router.patch("/settings/keys/{key_id}/toggle")
def toggle_api_key_status(
    key_id: int, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    db_key = db.query(models.ApiKey).filter(
        models.ApiKey.id == key_id,
        models.ApiKey.user_id == current_user.id
    ).first()
    if not db_key:
        raise HTTPException(status_code=404, detail="API Key not found")
        
    db_key.is_active = not db_key.is_active
    db.commit()
    db.refresh(db_key)
    return {
        "id": db_key.id,
        "provider": db_key.provider,
        "label": db_key.label,
        "is_active": db_key.is_active
    }

@router.delete("/settings/keys/{key_id}")
def delete_api_key_by_id(
    key_id: int, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    db_key = db.query(models.ApiKey).filter(
        models.ApiKey.id == key_id,
        models.ApiKey.user_id == current_user.id
    ).first()
    if not db_key:
        raise HTTPException(status_code=404, detail="API Key not found")
        
    db.delete(db_key)
    db.commit()
    return {"message": "API key deleted successfully"}

# Legacy endpoints compatibility - scoped to current_user
@router.get("/settings/api-key")
def get_api_key(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    first_key = db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "gemini", 
        models.ApiKey.is_active == True
    ).first()
    if first_key:
        return {"has_key": True, "masked_key": mask_api_key(decrypt_secret(first_key.key))}
    return {"has_key": False, "masked_key": ""}

class ApiKeyUpdate(BaseModel):
    api_key: str

@router.post("/settings/api-key")
def set_api_key(
    data: ApiKeyUpdate, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    api_key = data.api_key.strip()
    validate_key_with_provider("gemini", api_key)
    first_key = db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "gemini"
    ).first()
    if first_key:
        first_key.key = encrypt_secret(api_key)
        first_key.is_active = True
    else:
        first_key = models.ApiKey(
            user_id=current_user.id,
            provider="gemini", 
            key=encrypt_secret(api_key), 
            label="Gemini Key 1", 
            is_active=True
        )
        db.add(first_key)
    db.commit()
    return {"success": True}

@router.delete("/settings/api-key")
def delete_api_key(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "gemini"
    ).delete()
    db.commit()
    return {"success": True}

@router.get("/settings/openai-key")
def get_openai_key(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    first_key = db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "openai", 
        models.ApiKey.is_active == True
    ).first()
    if first_key:
        return {"has_key": True, "masked_key": mask_api_key(decrypt_secret(first_key.key))}
    return {"has_key": False, "masked_key": ""}

@router.post("/settings/openai-key")
def set_openai_key(
    data: ApiKeyUpdate, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    api_key = data.api_key.strip()
    validate_key_with_provider("openai", api_key)
    first_key = db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "openai"
    ).first()
    if first_key:
        first_key.key = encrypt_secret(api_key)
        first_key.is_active = True
    else:
        first_key = models.ApiKey(
            user_id=current_user.id,
            provider="openai", 
            key=encrypt_secret(api_key), 
            label="OpenAI Key 1", 
            is_active=True
        )
        db.add(first_key)
    db.commit()
    return {"success": True}

@router.delete("/settings/openai-key")
def delete_openai_key(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "openai"
    ).delete()
    db.commit()
    return {"success": True}

@router.get("/settings/groq-key")
def get_groq_key(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    first_key = db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "groq", 
        models.ApiKey.is_active == True
    ).first()
    if first_key:
        return {"has_key": True, "masked_key": mask_api_key(decrypt_secret(first_key.key))}
    return {"has_key": False, "masked_key": ""}

@router.post("/settings/groq-key")
def set_groq_key(
    data: ApiKeyUpdate, 
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    api_key = data.api_key.strip()
    validate_key_with_provider("groq", api_key)
    first_key = db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "groq"
    ).first()
    if first_key:
        first_key.key = encrypt_secret(api_key)
        first_key.is_active = True
    else:
        first_key = models.ApiKey(
            user_id=current_user.id,
            provider="groq", 
            key=encrypt_secret(api_key), 
            label="Groq Key 1", 
            is_active=True
        )
        db.add(first_key)
    db.commit()
    return {"success": True}

@router.delete("/settings/groq-key")
def delete_groq_key(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    db.query(models.ApiKey).filter(
        models.ApiKey.user_id == current_user.id,
        models.ApiKey.provider == "groq"
    ).delete()
    db.commit()
    return {"success": True}

@router.get("/download/desktop")
def download_desktop_app():
    from fastapi.responses import RedirectResponse
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    exe_candidates = [
        os.path.join(base_dir, "dist", "ERytmo_V2.exe"),
        os.path.join(base_dir, "uploads", "releases", "ERytmo_V2.exe"),
    ]
    for exe_path in exe_candidates:
        if os.path.exists(exe_path) and os.path.isfile(exe_path):
            return FileResponse(
                exe_path,
                filename="ERytmo_V2.exe",
                media_type="application/vnd.microsoft.portable-executable"
            )

    return RedirectResponse(
        url="https://github.com/sayfbel/ERytmo-Script-Converter/releases/latest/download/ERytmo_V2.exe",
        status_code=307
    )



