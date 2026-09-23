import os
import traceback
import logging
import json
import docx
from pypdf import PdfReader
from pydantic import BaseModel, Field
from google import genai
from google.genai import types

from exceptions import (
    ScriptImportError, UnsupportedFormatError, EmptyFileError,
    FileAccessError, FileCorruptedError, NoTimecodesFoundError
)
from script_validator import ScriptValidator, ValidationStatus, ValidationReport

# Configure conversion debug logger
LOG_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "conversion_debug.log")
logging.basicConfig(
    filename=LOG_FILE,
    level=logging.DEBUG,
    format="%(asctime)s [%(levelname)s] %(message)s"
)

def log_debug(message):
    try:
        logging.debug(message)
    except Exception:
        pass

def parse_json_response(text):
    if not text:
        return {}
    raw_response = text.strip()
    if raw_response.startswith("```json"):
        raw_response = raw_response[7:]
    elif raw_response.startswith("```"):
        raw_response = raw_response[3:]
    if raw_response.endswith("```"):
        raw_response = raw_response[:-3]
    raw_response = raw_response.strip()
    return json.loads(raw_response)

class ScriptCue(BaseModel):
    in_time: str = Field(alias="in", description="Timecode IN (e.g. 10:00:00:00)")
    out_time: str = Field(alias="out", description="Timecode OUT, empty if not available")
    character: str = Field(description="Character name speaking")
    dialogue: str = Field(description="Dialogue text")

class ScriptCues(BaseModel):
    cues: list[ScriptCue]

def get_gemini_api_key():
    key = os.environ.get("GEMINI_API_KEY")
    if key: return key
    config_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "config.json")
    if os.path.exists(config_path):
        try:
            with open(config_path, "r") as f:
                data = json.load(f)
                return data.get("GEMINI_API_KEY")
        except:
            pass
    return None

def parse_with_gemini(raw_text):
    api_key = get_gemini_api_key()
    if not api_key:
        log_debug("Gemini API key not found. Cannot perform AI parsing.")
        raise ScriptImportError(
            "Gemini API Key Missing",
            "To use the Universal Script Parser, you must provide a Gemini API Key in the settings.",
            "Please go to the Settings tab, enter your Gemini API Key, and try again."
        )
    
    try:
        client = genai.Client(api_key=api_key)
        prompt = f"""
You are an expert dubbing and adaptation script parser. Your task is to extract dialogue cues from the following raw script text.
The script may be in any language and any unstructured format. It may contain scene descriptions, action lines, and other noise.

For each dialogue cue you find, extract:
- 'in': The starting timecode (if any).
- 'out': The ending timecode (if any).
- 'character': The character speaking. Ignore scene directions.
- 'dialogue': The actual dialogue text spoken.

RULES:
- If a timecode is missing, leave it as an empty string ("").
- Do NOT include scene descriptions or action lines as dialogue or characters.
- Clean up character names (e.g., remove parentheticals like (V.O.) or (O.S.) or (ON THE PHONE)).
- Return ONLY valid JSON matching the requested schema.

SCRIPT TEXT (May be long):
{raw_text}
"""
        log_debug("Sending prompt to Gemini...")
        response = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=ScriptCues,
                temperature=0.1,
            ),
        )
        if response.text:
            data = parse_json_response(response.text)
            cues = data.get("cues", [])
            raw_rows = []
            for cue in cues:
                raw_rows.append({
                    "in": cue.get("in", ""),
                    "out": cue.get("out", ""),
                    "character": cue.get("character", "CHARACTER").strip(),
                    "dialogue": cue.get("dialogue", " ").strip()
                })
            log_debug(f"Gemini extracted {len(raw_rows)} cues successfully.")
            return raw_rows
        return []
    except Exception as e:
        log_debug(f"Gemini API parsing failed: {e}")
        raise ScriptImportError(
            "AI Parsing Failed",
            f"An error occurred while communicating with the Gemini API: {str(e)}",
            "Please check your internet connection or verify that your API key is valid."
        )

def check_chronological_order(cues):
    def tc_to_frames(tc):
        if not tc: return 0
        parts = tc.split(':')
        if len(parts) >= 3:
            h = int(parts[0])
            m = int(parts[1])
            s = int(parts[2])
            f = int(parts[3]) if len(parts) > 3 else 0
            return (h * 3600 + m * 60 + s) * 25 + f
        return 0

    def frames_to_tc(frames):
        f = frames % 25
        s = (frames // 25) % 60
        m = (frames // (25 * 60)) % 60
        h = (frames // (25 * 3600))
        return f"{h:02d}:{m:02d}:{s:02d}:{f:02d}"

    last_frames = 0
    for cue in cues:
        in_frames = tc_to_frames(cue["in"])
        if in_frames > 0:
            if in_frames < last_frames:
                cue["in"] = frames_to_tc(last_frames)
                in_frames = last_frames
            last_frames = in_frames
        
        out_frames = tc_to_frames(cue["out"])
        if out_frames > 0:
            if out_frames < in_frames:
                cue["out"] = frames_to_tc(in_frames)
            last_frames = tc_to_frames(cue["out"])

    return cues

def extract_audio_if_video(media_path):
    """
    Attempts to extract a lightweight MP3 audio from a video using ffmpeg.
    If it's already audio or ffmpeg is not installed, returns the original path.
    """
    ext = os.path.splitext(media_path)[1].lower()
    if ext in [".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"]:
        return media_path, False # Already audio, no temp file created

    import subprocess
    import tempfile
    
    temp_audio_path = os.path.join(tempfile.gettempdir(), f"extracted_audio_{os.path.basename(media_path)}.mp3")
    
    try:
        log_debug(f"Attempting to extract audio from video using ffmpeg: {media_path}")
        # Run ffmpeg to extract audio (no video, mp3 format, moderate quality)
        subprocess.run(
            ["ffmpeg", "-y", "-i", media_path, "-vn", "-c:a", "libmp3lame", "-q:a", "5", temp_audio_path],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL
        )
        log_debug(f"Successfully extracted audio to {temp_audio_path}")
        return temp_audio_path, True
    except (subprocess.CalledProcessError, FileNotFoundError) as e:
        log_debug(f"FFMPEG extraction failed or ffmpeg not found ({e}). Falling back to uploading original media.")
        return media_path, False

def align_timecodes_with_gemini(script_cues, media_path, start_timecode="00:00:00:00"):
    api_key = get_gemini_api_key()
    if not api_key:
        raise ScriptImportError(
            "Gemini API Key Missing",
            "To use the Universal Script Parser and Alignment, you must provide a Gemini API Key in the settings.",
            "Please go to the Settings tab, enter your Gemini API Key, and try again."
        )
    
    upload_path, is_temp = extract_audio_if_video(media_path)
    
    try:
        client = genai.Client(api_key=api_key)
        
        log_debug(f"Uploading media file to Gemini: {upload_path}")
        uploaded_media = client.files.upload(file=upload_path)
        
        log_debug("Waiting for file processing...")
        import time
        while uploaded_media.state.name == "PROCESSING":
            time.sleep(2)
            uploaded_media = client.files.get(name=uploaded_media.name)

        if uploaded_media.state.name == "FAILED":
             raise Exception("Gemini failed to process the media file.")
        
        cues_json = json.dumps(script_cues, ensure_ascii=False)

        prompt = f"""
You are an expert dubbing and adaptation script aligner. 
Your task is to take the provided JSON array of dialogue cues and accurately align them with the provided audio/video file.

RULES:
1. The media starts exactly at timecode: {start_timecode}.
2. For each cue in the JSON, find the exact 'in' (start) and 'out' (end) timestamps in SMPTE format (HH:MM:SS:FF) matching the dialogue spoken in the media. Assume 25 FPS.
3. If a cue already has an 'in' or 'out' timecode, DO NOT deviate massively from it. Just adjust it slightly to perfectly match the lip-sync or audio start.
4. CHRONOLOGICAL ORDER IS STRICTLY ENFORCED.
5. Return the exactly identical cues, but with the 'in' and 'out' fields corrected/generated. Do not modify the 'character' or 'dialogue' fields.
6. YOU MUST RETURN EXACTLY ONE JSON OBJECT PER LINE (JSONL format). Do NOT wrap in an array bracket [ or ]. Do NOT use markdown code blocks. Each line must be a valid JSON object representing one cue.

JSON CUES:
{cues_json}
"""
        log_debug("Sending prompt and media to Gemini (Streaming JSONL)...")
        response_stream = client.models.generate_content_stream(
            model='gemini-2.5-flash',
            contents=[uploaded_media, prompt],
            config=types.GenerateContentConfig(
                response_mime_type="text/plain",
                temperature=0.1,
            ),
        )
        
        def chronological_generator():
            buffer = ""
            last_frames = 0
            
            def tc_to_frames(tc):
                if not tc: return 0
                parts = tc.split(':')
                if len(parts) >= 3:
                    h = int(parts[0])
                    m = int(parts[1])
                    s = int(parts[2])
                    f = int(parts[3]) if len(parts) > 3 else 0
                    return (h * 3600 + m * 60 + s) * 25 + f
                return 0

            def frames_to_tc(frames):
                f = frames % 25
                s = (frames // 25) % 60
                m = (frames // (25 * 60)) % 60
                h = (frames // (25 * 3600))
                return f"{h:02d}:{m:02d}:{s:02d}:{f:02d}"

            try:
                for chunk in response_stream:
                    if chunk.text:
                        buffer += chunk.text
                        while "\n" in buffer:
                            line, buffer = buffer.split("\n", 1)
                            line = line.strip()
                            if not line or line.startswith("```") or line == "[" or line == "]": continue
                            if line.endswith(","): line = line[:-1]
                            try:
                                cue = json.loads(line)
                                out_cue = {
                                    "in": cue.get("in", ""),
                                    "out": cue.get("out", ""),
                                    "character": cue.get("character", "CHARACTER").strip(),
                                    "dialogue": cue.get("dialogue", " ").strip()
                                }
                                in_frames = tc_to_frames(out_cue["in"])
                                if in_frames > 0:
                                    if in_frames < last_frames:
                                        out_cue["in"] = frames_to_tc(last_frames)
                                        in_frames = last_frames
                                    last_frames = in_frames
                                
                                out_frames = tc_to_frames(out_cue["out"])
                                if out_frames > 0:
                                    if out_frames < in_frames:
                                        out_cue["out"] = frames_to_tc(in_frames)
                                    last_frames = tc_to_frames(out_cue["out"])
                                    
                                yield out_cue
                            except json.JSONDecodeError:
                                log_debug(f"Streaming JSON parse error on line: {line}")
                                continue

                # Process remaining buffer
                line = buffer.strip()
                if line and not line.startswith("```") and line != "[" and line != "]":
                    if line.endswith(","): line = line[:-1]
                    try:
                        cue = json.loads(line)
                        out_cue = {
                            "in": cue.get("in", ""),
                            "out": cue.get("out", ""),
                            "character": cue.get("character", "CHARACTER").strip(),
                            "dialogue": cue.get("dialogue", " ").strip()
                        }
                        in_frames = tc_to_frames(out_cue["in"])
                        if in_frames > 0:
                            if in_frames < last_frames:
                                out_cue["in"] = frames_to_tc(last_frames)
                                in_frames = last_frames
                            last_frames = in_frames
                        
                        out_frames = tc_to_frames(out_cue["out"])
                        if out_frames > 0:
                            if out_frames < in_frames:
                                out_cue["out"] = frames_to_tc(in_frames)
                            last_frames = tc_to_frames(out_cue["out"])
                        yield out_cue
                    except:
                        pass
            finally:
                try:
                     client.files.delete(name=uploaded_media.name)
                     log_debug("Deleted media file from Gemini servers.")
                except Exception as e:
                     log_debug(f"Failed to delete media file: {e}")
                     
                if is_temp and os.path.exists(upload_path):
                     try:
                         os.remove(upload_path)
                         log_debug("Deleted local temporary extracted audio file.")
                     except:
                         pass

        return chronological_generator()
    except Exception as e:
        log_debug(f"Gemini API Alignment failed: {e}")
        raise ScriptImportError(
            "AI Alignment Failed",
            f"An error occurred while communicating with the Gemini API or uploading the media: {str(e)}",
            "Please check your internet connection or verify that the media file is not corrupted."
        )

def extract_raw_text_docx(file_path):
    log_debug(f"Extracting raw text from DOCX: {file_path}")
    doc = docx.Document(file_path)
    text = []
    # Read paragraphs
    for p in doc.paragraphs:
        if p.text.strip(): text.append(p.text.strip())
    # Read tables
    for t in doc.tables:
        for r in t.rows:
            row_text = " | ".join([c.text.strip() for c in r.cells if c.text.strip()])
            if row_text: text.append(row_text)
    return "\n".join(text)

def extract_raw_text_pdf(file_path):
    log_debug(f"Extracting raw text from PDF: {file_path}")
    reader = PdfReader(file_path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def extract_raw_text_txt(file_path):
    log_debug(f"Extracting raw text from TXT: {file_path}")
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def safe_validate_and_convert(input_path, output_docx_path=None, export_mode="3line"):
    """
    Safely validates, extracts raw text, and delegates parsing to Gemini AI.
    Never crashes. Returns (report, raw_rows, format_a_cues).
    export_mode: '3line' (IN only), '4line' (IN & OUT separate lines), 'inline' (IN - OUT combined line)
    """
    log_debug(f"--- Starting Validation & AI Import for: {input_path} ---")
    
    # 1. Run Pre-Validation
    report = ScriptValidator.validate_file(input_path)
    
    if report.status == ValidationStatus.INVALID:
        log_debug(f"Validation Failed: {report.user_title} - {report.user_message}")
        return report, [], []

    # 2. Extract Raw Text
    ext = os.path.splitext(input_path)[1].lower()
    raw_text = ""
    try:
        if ext == ".docx":
            raw_text = extract_raw_text_docx(input_path)
        elif ext == ".pdf":
            raw_text = extract_raw_text_pdf(input_path)
        elif ext in [".txt", ".text"]:
            raw_text = extract_raw_text_txt(input_path)
        else:
            raise UnsupportedFormatError(ext)

        if not raw_text.strip():
            raise EmptyFileError(os.path.basename(input_path))

        # 3. Process with AI
        raw_rows = parse_with_gemini(raw_text)

        if not raw_rows:
            raise NoTimecodesFoundError(os.path.basename(input_path))

        report.status = ValidationStatus.NORMALIZABLE
        report.user_title = "Script Parsed Successfully with AI"
        report.user_message = f"The script was successfully extracted using the AI Parser. Detected {len(raw_rows)} cues."

        format_a_cues = []
        for r in raw_rows:
            format_a_cues.append({
                "in": r["in"],
                "out": r.get("out", ""),
                "character": r["character"],
                "dialogue": r["dialogue"]
            })

        # Save to output docx if requested
        if output_docx_path:
            out_doc = docx.Document()
            for idx, cue in enumerate(format_a_cues):
                tc_in = cue["in"]
                tc_out = cue.get("out", "")
                
                if export_mode == "4line" and tc_out:
                    out_doc.add_paragraph(tc_in)
                    out_doc.add_paragraph(tc_out)
                elif export_mode == "inline" and tc_out:
                    out_doc.add_paragraph(f"{tc_in} - {tc_out}")
                else: # '3line' default
                    if tc_in:
                        out_doc.add_paragraph(tc_in)

                out_doc.add_paragraph(cue["character"])
                out_doc.add_paragraph(cue["dialogue"])
                
                if idx < len(format_a_cues) - 1:
                    out_doc.add_paragraph("") # Blank line separator
            out_doc.save(output_docx_path)
            log_debug(f"Saved Format DOCX (export_mode={export_mode}) to: {output_docx_path}")

        log_debug(f"Successfully processed {len(raw_rows)} cues. Status: {report.status}")
        return report, raw_rows, format_a_cues

    except ScriptImportError as e:
        tb = traceback.format_exc()
        log_debug(f"ScriptImportError Caught:\n{tb}")
        report.status = ValidationStatus.INVALID
        report.user_title = e.user_title
        report.user_message = e.user_message
        report.suggestion = e.suggestion
        report.diagnostic_info += f"\n\nException:\n{tb}"
        return report, [], []

    except Exception as e:
        tb = traceback.format_exc()
        log_debug(f"Unexpected Exception Caught:\n{tb}")
        report.status = ValidationStatus.INVALID
        report.user_title = "Unable to Understand Script Format"
        report.user_message = "An unexpected error occurred while parsing the script content."
        report.suggestion = "The structure of this script might be entirely corrupted or there is a bug in the AI integration. Please contact Support."
        report.diagnostic_info += f"\n\nUnexpected Exception:\n{tb}"
        return report, [], []

def extract_and_convert(input_path, output_docx_path=None, export_mode="3line"):
    """
    Helper function for direct conversion.
    Returns (raw_rows, format_a_cues).
    """
    report, raw_rows, format_a_cues = safe_validate_and_convert(input_path, output_docx_path, export_mode=export_mode)
    if report.status == ValidationStatus.INVALID:
        raise ScriptImportError(report.user_title, report.user_message, report.suggestion, report.diagnostic_info)
    return raw_rows, format_a_cues
