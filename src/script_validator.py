import os
import re
import traceback
import platform
import docx
from pypdf import PdfReader
from exceptions import (
    ScriptImportError, UnsupportedFormatError, EmptyFileError,
    FileAccessError, FileCorruptedError, NoTimecodesFoundError, MissingRequiredFieldsError
)

class ValidationStatus:
    VALID = "VALID"
    NORMALIZABLE = "NORMALIZABLE"
    INVALID = "INVALID"

class ValidationReport:
    def __init__(self, status, user_title, user_message, suggestion=None, normalizations=None, issues=None, diagnostic_info=None):
        self.status = status
        self.user_title = user_title
        self.user_message = user_message
        self.suggestion = suggestion or "Please check your script format and try importing it again."
        self.normalizations = normalizations or []
        self.issues = issues or []
        self.diagnostic_info = diagnostic_info or ""

class ScriptValidator:
    SUPPORTED_EXTENSIONS = [".docx", ".pdf", ".txt", ".text"]

    @classmethod
    def validate_file(cls, file_path):
        """
        Pre-validates script file before parsing.
        Returns a ValidationReport object.
        """
        diagnostic_lines = [
            "=== ERytmo Script Converter Diagnostic Report ===",
            f"OS: {platform.system()} {platform.release()}",
            f"Python: {platform.python_version()}",
            f"File Path: {file_path}",
        ]

        if not file_path or not os.path.exists(file_path):
            diag = "\n".join(diagnostic_lines + ["Error: File path does not exist."])
            return ValidationReport(
                ValidationStatus.INVALID,
                "File Not Found",
                "The specified script file could not be found on disk.",
                "Please verify the file path and select a valid script file.",
                diagnostic_info=diag
            )

        file_name = os.path.basename(file_path)
        ext = os.path.splitext(file_path)[1].lower()
        file_size = os.path.getsize(file_path)
        
        diagnostic_lines.append(f"File Size: {file_size} bytes")
        diagnostic_lines.append(f"File Extension: {ext}")

        # Check 1: Supported Format Extension
        if ext not in cls.SUPPORTED_EXTENSIONS:
            diag = "\n".join(diagnostic_lines + [f"Error: Unsupported format {ext}"])
            return ValidationReport(
                ValidationStatus.INVALID,
                "Unsupported File Format",
                f"The file extension '{ext}' is currently not supported.",
                "Please select a supported script file format (.docx, .pdf, .txt, or .text).",
                diagnostic_info=diag
            )

        # Check 2: Empty File Check
        if file_size == 0:
            diag = "\n".join(diagnostic_lines + ["Error: File size is 0 bytes."])
            return ValidationReport(
                ValidationStatus.INVALID,
                "Empty Script File",
                f"The selected file '{file_name}' contains no content or data (0 bytes).",
                "Please select a non-empty script file containing dialogue cues and timecodes.",
                diagnostic_info=diag
            )

        # Check 3: File Read / Corruption Check
        try:
            if ext == ".docx":
                doc = docx.Document(file_path)
                has_tables = len(doc.tables) > 0
                has_paragraphs = len(doc.paragraphs) > 0
                diagnostic_lines.append(f"DOCX Structure: Tables={len(doc.tables)}, Paragraphs={len(doc.paragraphs)}")
                if not has_tables and not has_paragraphs:
                    diag = "\n".join(diagnostic_lines + ["Error: DOCX document contains no tables or paragraphs."])
                    return ValidationReport(
                        ValidationStatus.INVALID,
                        "Empty Document Content",
                        f"The document '{file_name}' contains no readable tables or paragraphs.",
                        "Please open the file in Word and ensure it contains script content.",
                        diagnostic_info=diag
                    )

            elif ext == ".pdf":
                reader = PdfReader(file_path)
                diagnostic_lines.append(f"PDF Structure: Pages={len(reader.pages)}")
                if len(reader.pages) == 0:
                    diag = "\n".join(diagnostic_lines + ["Error: PDF has 0 pages."])
                    return ValidationReport(
                        ValidationStatus.INVALID,
                        "Empty PDF Document",
                        f"The PDF file '{file_name}' contains no pages.",
                        "Please check the PDF file and try again.",
                        diagnostic_info=diag
                    )

            elif ext in [".txt", ".text"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    lines = f.readlines()
                diagnostic_lines.append(f"Text Lines: {len(lines)}")
                if not lines:
                    diag = "\n".join(diagnostic_lines + ["Error: Text file is empty."])
                    return ValidationReport(
                        ValidationStatus.INVALID,
                        "Empty Text File",
                        f"The text file '{file_name}' contains no lines.",
                        "Please select a valid script file.",
                        diagnostic_info=diag
                    )

        except Exception as e:
            tb = traceback.format_exc()
            diag = "\n".join(diagnostic_lines + [f"Read Exception:\n{tb}"])
            return ValidationReport(
                ValidationStatus.INVALID,
                "Unable to Read File",
                f"Could not read or parse '{file_name}'. The file may be corrupt or open in another application.",
                "Ensure the file is not locked by another program and is a valid document.",
                diagnostic_info=diag
            )

        # Check 4: Deep Content & Structure Scan (Timecodes & Characters)
        report = cls._analyze_script_structure(file_path, ext, diagnostic_lines)
        return report

    @classmethod
    def _analyze_script_structure(cls, file_path, ext, diagnostic_lines):
        normalizations = []
        issues = []
        found_timecodes = 0

        if ext == ".docx":
            doc = docx.Document(file_path)
            if doc.tables:
                table = doc.tables[0]
                keywords = ['TIMECODE', 'TIME', 'CODE', 'IN', 'OUT', 'SHOT', 'CHARACTER', 'PERSO', 'PERSONNAGE', 'DIALOGUE', 'DIALOG', 'TITLE', 'SCENE', 'TEXT', 'SPEECH', 'SPEAKER', 'HORODATAGE']
                
                # Scan table rows for column header keywords
                best_row_idx = 0
                max_kw = 0
                for r_idx in range(min(6, len(table.rows))):
                    cells_text = [c.text.strip() for c in table.rows[r_idx].cells]
                    if any(re.search(r'\b\d{1,2}:\d{2}:\d{2}', c) for c in cells_text if c):
                        continue
                    found_kw = set()
                    for cell_text in cells_text:
                        norm_h = re.sub(r'[^A-Z0-9]', '', cell_text.upper())
                        for kw in keywords:
                            if kw in norm_h:
                                found_kw.add(kw)
                    if len(found_kw) > max_kw:
                        max_kw = len(found_kw)
                        best_row_idx = r_idx

                hdr_cells = [c.text.strip() for c in table.rows[best_row_idx].cells]
                diagnostic_lines.append(f"Detected Header Row Index: {best_row_idx}")
                diagnostic_lines.append(f"Header Columns: {hdr_cells}")

                if best_row_idx > 0:
                    normalizations.append(f"Detected multi-row document header; aligned column headers at row {best_row_idx + 1}.")

                # Count timecodes in data rows using flexible search
                for r_idx in range(best_row_idx + 1, len(table.rows)):
                    for cell in table.rows[r_idx].cells:
                        if re.search(r'\b\d{1,2}:\d{2}:\d{2}', cell.text.strip()):
                            found_timecodes += 1
                            break

                norm_hdrs = [re.sub(r'[^A-Z0-9]', '', h.upper()) for h in hdr_cells]
                if any(k in h for h in norm_hdrs for k in ["OUT", "STF", "END", "TCOUT"]):
                    normalizations.append("Detected OUT timecode column. Stripped OUT timecodes to prevent ERytmo listbox parsing issues.")

                if any(k in h for h in norm_hdrs for k in ["PERSO", "SPEAKER", "PERSONNAGE", "TITLE", "CHARACTER"]):
                    normalizations.append("Mapped custom column header to standard 'CHARACTER' field.")

            else:
                for p in doc.paragraphs:
                    if re.search(r'\b\d{1,2}:\d{2}:\d{2}', p.text.strip()):
                        found_timecodes += 1

        elif ext == ".pdf":
            reader = PdfReader(file_path)
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    for line in txt.split("\n"):
                        if re.search(r'\b\d{1,2}:\d{2}:\d{2}', line.strip()):
                            found_timecodes += 1

        elif ext in [".txt", ".text"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                for line in f:
                    if re.search(r'\b\d{1,2}:\d{2}:\d{2}', line.strip()):
                        found_timecodes += 1

        diagnostic_lines.append(f"Total Timecodes Found: {found_timecodes}")

        # Evaluate Result
        if found_timecodes == 0:
            diag = "\n".join(diagnostic_lines + ["Error: 0 timecodes found matching HH:MM:SS:FF or HH:MM:SS pattern."])
            return ValidationReport(
                ValidationStatus.INVALID,
                "Unable to Understand Script Format",
                "The structure of this script is currently not supported because no valid timecodes (e.g. 05:00:00:00) were detected.",
                "Please check your script format and ensure timecodes are formatted as HH:MM:SS:FF or HH:MM:SS, then try importing again.",
                diagnostic_info=diag
            )

        if normalizations:
            status = ValidationStatus.NORMALIZABLE
            user_title = "Script Format Detected & Auto-Normalized"
            user_message = f"The script was successfully validated with {found_timecodes} dialogue cues detected. Minor formatting variations were automatically normalized."
        else:
            status = ValidationStatus.VALID
            user_title = "Script Format Valid"
            user_message = f"The script was successfully validated. {found_timecodes} dialogue cues are perfectly formatted."

        diag = "\n".join(diagnostic_lines + ["Validation Status: SUCCESS"])
        return ValidationReport(
            status,
            user_title,
            user_message,
            normalizations=normalizations,
            issues=issues,
            diagnostic_info=diag
        )
