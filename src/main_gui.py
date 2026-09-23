import os
import sys
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QFileDialog, QTableWidget, QTableWidgetItem, QHeaderView,
    QProgressBar, QFrame, QSplitter, QMessageBox, QDialog, QTextEdit, QLineEdit,
    QComboBox
)
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtGui import QIcon, QPixmap, QFont, QColor

# Import script parser and validator
from script_parser import safe_validate_and_convert, extract_and_convert, align_timecodes_with_gemini
from script_validator import ValidationStatus, ValidationReport

class ParseThread(QThread):
    finished_signal = Signal(object, object, object)

    def __init__(self, input_path):
        super().__init__()
        self.input_path = input_path

    def run(self):
        report, raw_rows, format_a_cues = safe_validate_and_convert(self.input_path)
        self.finished_signal.emit(report, raw_rows, format_a_cues)

class AlignThread(QThread):
    row_updated_signal = Signal(int, dict)
    finished_signal = Signal(object)

    def __init__(self, cues, media_path, start_tc):
        super().__init__()
        self.cues = cues
        self.media_path = media_path
        self.start_tc = start_tc

    def run(self):
        try:
            generator = align_timecodes_with_gemini(self.cues, self.media_path, self.start_tc)
            aligned_cues = []
            for idx, cue in enumerate(generator):
                aligned_cues.append(cue)
                self.row_updated_signal.emit(idx, cue)
            self.finished_signal.emit(aligned_cues)
        except Exception as e:
            self.finished_signal.emit(e)

class DropAreaWidget(QFrame):
    file_dropped_signal = Signal(str)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.setObjectName("DropArea")

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        urls = event.mimeData().urls()
        if urls:
            file_path = urls[0].toLocalFile()
            ext = os.path.splitext(file_path)[1].lower()
            if ext in [".docx", ".pdf", ".txt", ".text"]:
                self.file_dropped_signal.emit(file_path)

class SupportDialog(QDialog):
    def __init__(self, report, parent=None):
        super().__init__(parent)
        self.report = report
        self.setWindowTitle("Import Support & Technical Diagnostics")
        self.resize(600, 450)
        self.setStyleSheet("""
            QDialog { background-color: #FFFFFF; }
            QLabel { font-family: 'Segoe UI', Arial, sans-serif; color: #0F172A; }
            QTextEdit { background-color: #F8FAFC; border: 1px solid #CBD5E1; border-radius: 6px; font-family: 'Consolas', monospace; font-size: 11px; color: #334155; }
            QPushButton#CopyBtn { background-color: #2563EB; color: #FFFFFF; font-weight: bold; border-radius: 6px; padding: 8px 16px; border: none; }
            QPushButton#CopyBtn:hover { background-color: #1D4ED8; }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)

        title_lbl = QLabel("Technical Diagnostic Report")
        title_lbl.setFont(QFont("Segoe UI", 12, QFont.Bold))
        layout.addWidget(title_lbl)

        desc_lbl = QLabel("You can copy the diagnostic information below and send it to Support to request format compatibility.")
        desc_lbl.setFont(QFont("Segoe UI", 9))
        desc_lbl.setStyleSheet("color: #64748B;")
        desc_lbl.setWordWrap(True)
        layout.addWidget(desc_lbl)
        layout.addSpacing(10)

        self.diag_edit = QTextEdit()
        self.diag_edit.setReadOnly(True)
        if isinstance(report, ValidationReport):
             self.diag_edit.setText(report.diagnostic_info)
        else:
             self.diag_edit.setText(str(report))
        layout.addWidget(self.diag_edit)
        layout.addSpacing(10)

        btn_box = QHBoxLayout()
        copy_btn = QPushButton("📋 Copy Diagnostic Info to Clipboard")
        copy_btn.setObjectName("CopyBtn")
        copy_btn.clicked.connect(self.copy_to_clipboard)

        close_btn = QPushButton("Close")
        close_btn.clicked.connect(self.accept)

        btn_box.addWidget(copy_btn)
        btn_box.addStretch()
        btn_box.addWidget(close_btn)
        layout.addLayout(btn_box)

    def copy_to_clipboard(self):
        clipboard = QApplication.clipboard()
        clipboard.setText(self.diag_edit.toPlainText())
        QMessageBox.information(self, "Copied", "Diagnostic information copied to clipboard!")

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.selected_file_path = None
        self.loaded_media_path = None
        self.raw_rows = []
        self.format_a_cues = []
        self.last_report = None

        self.setWindowTitle("ERytmo Script Converter v2 - AI Alignment")
        self.resize(1120, 750)
        self.setMinimumSize(950, 650)

        base_dir = os.path.dirname(os.path.abspath(__file__))
        logo_png = os.path.join(base_dir, "app_logo.png")
        logo_ico = os.path.join(base_dir, "app_logo.ico")
        if os.path.exists(logo_ico):
            self.setWindowIcon(QIcon(logo_ico))
        elif os.path.exists(logo_png):
            self.setWindowIcon(QIcon(logo_png))

        self.init_ui(logo_png)

    def init_ui(self, logo_path):
        self.setStyleSheet("""
            QMainWindow { background-color: #F8FAFC; }
            QWidget { font-family: 'Segoe UI', Arial, sans-serif; color: #0F172A; }
            QFrame#HeaderFrame { background-color: #FFFFFF; border-bottom: 1px solid #E2E8F0; }
            QFrame#DropArea { background-color: #FFFFFF; border: 2px dashed #CBD5E1; border-radius: 12px; }
            QFrame#DropArea:hover { border-color: #2563EB; background-color: #F1F5F9; }
            QFrame#CardFrame { background-color: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 10px; }
            QPushButton#PrimaryButton { background-color: #2563EB; color: #FFFFFF; font-size: 13px; font-weight: bold; border-radius: 6px; padding: 8px 16px; border: none; }
            QPushButton#PrimaryButton:hover { background-color: #1D4ED8; }
            QPushButton#PrimaryButton:disabled { background-color: #94A3B8; }
            QPushButton#DownloadButton { background-color: #10B981; color: #FFFFFF; font-size: 14px; font-weight: bold; border-radius: 8px; padding: 10px 22px; border: none; }
            QPushButton#DownloadButton:hover { background-color: #059669; }
            QPushButton#DownloadButton:disabled { background-color: #E2E8F0; color: #94A3B8; }
            QPushButton#SupportButton { background-color: #EF4444; color: #FFFFFF; font-size: 12px; font-weight: bold; border-radius: 6px; padding: 6px 12px; border: none; }
            QPushButton#SupportButton:hover { background-color: #DC2626; }
            QPushButton#SecondaryButton { background-color: #FFFFFF; color: #334155; font-size: 13px; font-weight: 600; border-radius: 6px; padding: 8px 16px; border: 1px solid #CBD5E1; }
            QPushButton#SecondaryButton:hover { background-color: #F1F5F9; }
            QLineEdit { border: 1px solid #CBD5E1; border-radius: 4px; padding: 4px; background: #FFFFFF; }
            QTableWidget { background-color: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 8px; gridline-color: #F1F5F9; font-size: 12px; color: #0F172A; }
            QTableWidget::item { padding: 6px; }
            QHeaderView::section { background-color: #F1F5F9; color: #334155; font-weight: bold; border: none; border-bottom: 1px solid #CBD5E1; padding: 8px; }
            QProgressBar { border: none; background-color: #E2E8F0; height: 4px; border-radius: 2px; }
            QProgressBar::chunk { background-color: #2563EB; }
        """)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # Header
        header_frame = QFrame()
        header_frame.setObjectName("HeaderFrame")
        header_layout = QHBoxLayout(header_frame)
        header_layout.setContentsMargins(20, 14, 20, 14)

        if os.path.exists(logo_path):
            logo_label = QLabel()
            pix = QPixmap(logo_path).scaled(40, 40, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            logo_label.setPixmap(pix)
            header_layout.addWidget(logo_label)

        title_box = QVBoxLayout()
        title_label = QLabel("ERytmo Script Converter v2")
        title_label.setFont(QFont("Segoe UI", 15, QFont.Bold))
        subtitle_label = QLabel("Convert Scripts and Align Timecodes with AI")
        subtitle_label.setStyleSheet("color: #64748B;")
        title_box.addWidget(title_label)
        title_box.addWidget(subtitle_label)
        header_layout.addLayout(title_box)
        header_layout.addStretch()
        main_layout.addWidget(header_frame)

        # Body
        body_widget = QWidget()
        body_layout = QVBoxLayout(body_widget)
        body_layout.setContentsMargins(20, 16, 20, 16)
        body_layout.setSpacing(12)

        # Row 1: Script Load
        top_row = QHBoxLayout()
        top_row.setSpacing(12)

        self.drop_area = DropAreaWidget()
        drop_layout = QVBoxLayout(self.drop_area)
        drop_text_label = QLabel("Drag & Drop Script File (.docx, .pdf, .txt)")
        drop_text_label.setFont(QFont("Segoe UI", 11, QFont.Bold))
        drop_text_label.setAlignment(Qt.AlignCenter)
        browse_btn = QPushButton("Browse File...")
        browse_btn.setObjectName("SecondaryButton")
        browse_btn.clicked.connect(self.on_browse_clicked)
        drop_layout.addWidget(drop_text_label)
        drop_layout.addWidget(browse_btn, 0, Qt.AlignCenter)
        self.drop_area.file_dropped_signal.connect(self.on_file_selected)
        top_row.addWidget(self.drop_area, 2)

        self.file_card = QFrame()
        self.file_card.setObjectName("CardFrame")
        card_layout = QVBoxLayout(self.file_card)
        self.file_name_label = QLabel("No File Loaded")
        self.file_name_label.setFont(QFont("Segoe UI", 11, QFont.Bold))
        self.status_badge_label = QLabel("Select or drop a script file to parse.")
        
        btn_layout = QHBoxLayout()
        self.export_combo = QComboBox()
        self.export_combo.addItems(["ERytmo Factory (Standard)", "Mosaic (Table)"])
        self.export_combo.setStyleSheet("border: 1px solid #CBD5E1; border-radius: 6px; padding: 6px; font-weight: bold; background: white;")
        self.export_combo.setEnabled(False)
        self.download_btn = QPushButton("📥 Download DOCX")
        self.download_btn.setObjectName("DownloadButton")
        self.download_btn.setEnabled(False)
        self.download_btn.clicked.connect(self.on_download_clicked)
        self.support_btn = QPushButton("🎧 Support")
        self.support_btn.setObjectName("SupportButton")
        self.support_btn.setVisible(False)
        self.support_btn.clicked.connect(self.on_support_clicked)
        btn_layout.addWidget(self.export_combo)
        btn_layout.addWidget(self.download_btn, 1)
        btn_layout.addWidget(self.support_btn)
        
        card_layout.addWidget(self.file_name_label)
        card_layout.addWidget(self.status_badge_label)
        card_layout.addStretch()
        card_layout.addLayout(btn_layout)
        top_row.addWidget(self.file_card, 1)

        body_layout.addLayout(top_row)

        # Row 2: Media Alignment
        media_row = QHBoxLayout()
        self.media_card = QFrame()
        self.media_card.setObjectName("CardFrame")
        media_layout = QVBoxLayout(self.media_card)
        
        media_title = QLabel("Optional: AI Media Timecode Verification")
        media_title.setFont(QFont("Segoe UI", 11, QFont.Bold))
        
        media_controls = QHBoxLayout()
        self.load_media_btn = QPushButton("🎵 Load Media (Video/Audio)")
        self.load_media_btn.setObjectName("SecondaryButton")
        self.load_media_btn.clicked.connect(self.on_load_media_clicked)
        
        self.media_path_label = QLabel("No media loaded")
        self.media_path_label.setStyleSheet("color: #64748B;")
        
        tc_label = QLabel("Start TC:")
        self.start_tc_input = QLineEdit("00:00:00:00")
        self.start_tc_input.setMaximumWidth(100)
        
        self.verify_btn = QPushButton("✨ Verify Timecodes (AI)")
        self.verify_btn.setObjectName("PrimaryButton")
        self.verify_btn.setEnabled(False)
        self.verify_btn.clicked.connect(self.on_verify_clicked)
        
        media_controls.addWidget(self.load_media_btn)
        media_controls.addWidget(self.media_path_label)
        media_controls.addStretch()
        media_controls.addWidget(tc_label)
        media_controls.addWidget(self.start_tc_input)
        media_controls.addWidget(self.verify_btn)
        
        media_layout.addWidget(media_title)
        media_layout.addLayout(media_controls)
        media_row.addWidget(self.media_card)
        body_layout.addLayout(media_row)

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setVisible(False)
        body_layout.addWidget(self.progress_bar)

        # Tables
        tables_splitter = QSplitter(Qt.Horizontal)
        t1_box = QWidget()
        t1_layout = QVBoxLayout(t1_box)
        t1_layout.setContentsMargins(0, 0, 0, 0)
        t1_label = QLabel("1. Original Imported Script (Raw)")
        t1_label.setFont(QFont("Segoe UI", 10, QFont.Bold))
        self.table_raw = QTableWidget(0, 4)
        self.table_raw.setHorizontalHeaderLabels(["IN", "OUT", "Character", "Dialogue / Text"])
        self.table_raw.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        t1_layout.addWidget(t1_label)
        t1_layout.addWidget(self.table_raw)

        t2_box = QWidget()
        t2_layout = QVBoxLayout(t2_box)
        t2_layout.setContentsMargins(0, 0, 0, 0)
        t2_label = QLabel("2. Converted / Aligned Script (Output Preview)")
        t2_label.setFont(QFont("Segoe UI", 10, QFont.Bold))
        t2_label.setStyleSheet("color: #166534;")
        self.table_converted = QTableWidget(0, 4)
        self.table_converted.setHorizontalHeaderLabels(["IN", "OUT", "Character", "Dialogue / Text"])
        self.table_converted.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        t2_layout.addWidget(t2_label)
        t2_layout.addWidget(self.table_converted)

        tables_splitter.addWidget(t1_box)
        tables_splitter.addWidget(t2_box)
        body_layout.addWidget(tables_splitter, 1)

        main_layout.addWidget(body_widget)

    def on_browse_clicked(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select Script File", "", "Script Files (*.docx *.pdf *.txt *.text)")
        if file_path:
            self.on_file_selected(file_path)

    def on_load_media_clicked(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select Media File", "", "Media Files (*.mp4 *.mkv *.avi *.mp3 *.wav *.m4a)")
        if file_path:
            self.loaded_media_path = file_path
            self.media_path_label.setText(os.path.basename(file_path))
            if self.raw_rows:
                self.verify_btn.setEnabled(True)

    def on_file_selected(self, file_path):
        self.selected_file_path = file_path
        self.file_name_label.setText(os.path.basename(file_path))
        self.status_badge_label.setText("Parsing script with AI...")
        self.progress_bar.setVisible(True)
        self.export_combo.setEnabled(False)
        self.download_btn.setEnabled(False)
        self.support_btn.setVisible(False)
        self.verify_btn.setEnabled(False)

        self.table_raw.setRowCount(0)
        self.table_converted.setRowCount(0)

        self.thread = ParseThread(file_path)
        self.thread.finished_signal.connect(self.on_parse_finished)
        self.thread.start()

    def on_verify_clicked(self):
        if not self.format_a_cues or not self.loaded_media_path: return
        self.status_badge_label.setText("Uploading media and aligning timecodes... (This may take a minute)")
        self.progress_bar.setVisible(True)
        self.verify_btn.setEnabled(False)
        
        self.align_thread = AlignThread(self.format_a_cues, self.loaded_media_path, self.start_tc_input.text())
        self.align_thread.row_updated_signal.connect(self.on_align_row_updated)
        self.align_thread.finished_signal.connect(self.on_align_finished)
        self.align_thread.start()

    def on_align_row_updated(self, r_idx, c):
        if r_idx < self.table_converted.rowCount():
            self.table_converted.setItem(r_idx, 0, QTableWidgetItem(c.get("in", "")))
            self.table_converted.setItem(r_idx, 1, QTableWidgetItem(c.get("out", "")))
            item_char = QTableWidgetItem(c.get("character", ""))
            item_char.setForeground(QColor("#166534"))
            item_char.setFont(QFont("Segoe UI", 9, QFont.Bold))
            self.table_converted.setItem(r_idx, 2, item_char)
            self.table_converted.setItem(r_idx, 3, QTableWidgetItem(c.get("dialogue", "")))
            self.table_converted.scrollToItem(self.table_converted.item(r_idx, 0))

    def on_align_finished(self, result):
        self.progress_bar.setVisible(False)
        self.verify_btn.setEnabled(True)
        if isinstance(result, Exception):
            QMessageBox.critical(self, "Alignment Failed", str(result))
            dlg = SupportDialog(result, self)
            dlg.exec()
            return
        
        self.format_a_cues = result
        self.status_badge_label.setText(f"🟢 Successfully aligned {len(self.format_a_cues)} cues to media!")
        self.populate_converted_table(self.format_a_cues)

    def populate_converted_table(self, cues):
        self.table_converted.setRowCount(len(cues))
        for r_idx, c in enumerate(cues):
            self.table_converted.setItem(r_idx, 0, QTableWidgetItem(c.get("in", "")))
            self.table_converted.setItem(r_idx, 1, QTableWidgetItem(c.get("out", "")))
            item_char = QTableWidgetItem(c.get("character", ""))
            item_char.setForeground(QColor("#166534"))
            item_char.setFont(QFont("Segoe UI", 9, QFont.Bold))
            self.table_converted.setItem(r_idx, 2, item_char)
            self.table_converted.setItem(r_idx, 3, QTableWidgetItem(c.get("dialogue", "")))

    def on_parse_finished(self, report, raw_rows, format_a_cues):
        self.progress_bar.setVisible(False)
        self.last_report = report

        if report.status == ValidationStatus.INVALID:
            self.status_badge_label.setText(f"🔴 <b>{report.user_title}</b><br>{report.user_message}")
            self.download_btn.setEnabled(False)
            self.support_btn.setVisible(True)
            QMessageBox.critical(self, report.user_title, f"{report.user_message}\n\nSuggestion: {report.suggestion}")
            return

        self.raw_rows = raw_rows
        self.format_a_cues = format_a_cues
        self.export_combo.setEnabled(True)
        self.download_btn.setEnabled(True)
        if self.loaded_media_path:
            self.verify_btn.setEnabled(True)

        self.status_badge_label.setText(f"🟢 Parsed {len(raw_rows)} dialogue cues.")

        self.table_raw.setRowCount(len(raw_rows))
        for r_idx, r in enumerate(raw_rows):
            self.table_raw.setItem(r_idx, 0, QTableWidgetItem(r.get("in", "")))
            self.table_raw.setItem(r_idx, 1, QTableWidgetItem(r.get("out", "")))
            self.table_raw.setItem(r_idx, 2, QTableWidgetItem(r.get("character", "")))
            self.table_raw.setItem(r_idx, 3, QTableWidgetItem(r.get("dialogue", "")))

        self.populate_converted_table(format_a_cues)

    def on_download_clicked(self):
        if not self.selected_file_path or not self.format_a_cues: return
        
        export_mode = self.export_combo.currentText()
        is_mosaic = "Mosaic" in export_mode
        
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        suffix = "Mosaic" if is_mosaic else "ERytmo"
        default_path = os.path.join(os.path.dirname(self.selected_file_path), f"{base_name}_{suffix}_Format.docx")
        
        output_path, _ = QFileDialog.getSaveFileName(self, "Save Converted DOCX", default_path, "Word Documents (*.docx)")
        if output_path:
            try:
                import docx
                out_doc = docx.Document()
                
                if is_mosaic:
                    table = out_doc.add_table(rows=1, cols=4)
                    table.style = 'Table Grid'
                    hdr_cells = table.rows[0].cells
                    hdr_cells[0].text = 'Timecode IN'
                    hdr_cells[1].text = 'Timecode OUT'
                    hdr_cells[2].text = 'Character'
                    hdr_cells[3].text = 'Dialogue / Text'
                    
                    for cue in self.format_a_cues:
                        row_cells = table.add_row().cells
                        row_cells[0].text = cue.get("in", "")
                        row_cells[1].text = cue.get("out", "")
                        row_cells[2].text = cue.get("character", "")
                        row_cells[3].text = cue.get("dialogue", "")
                else:
                    for cue in self.format_a_cues:
                        tc_in = cue.get("in", "")
                        if tc_in: out_doc.add_paragraph(tc_in)
                        out_doc.add_paragraph(cue.get("character", ""))
                        out_doc.add_paragraph(cue.get("dialogue", ""))
                        out_doc.add_paragraph("")
                        
                out_doc.save(output_path)
                QMessageBox.information(self, "Success", f"Saved to {output_path}")
            except Exception as e:
                QMessageBox.critical(self, "Error Saving File", str(e))

    def on_support_clicked(self):
        if self.last_report:
            dlg = SupportDialog(self.last_report, self)
            dlg.exec()

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
